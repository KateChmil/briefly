import io

import pytest
from docx import Document

from app.services.extract import extract_text


def test_extract_txt():
    assert extract_text("a.txt", b"hello\nworld") == "hello\nworld"


def test_extract_docx():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.save(buf)
    text = extract_text("notes.docx", buf.getvalue())
    assert "First paragraph" in text and "Second paragraph" in text


def test_extract_unsupported():
    with pytest.raises(ValueError):
        extract_text("img.png", b"\x89PNG")


def test_upload_source(client):
    sid = client.post("/api/spaces", json={"name": "Bio"}).json()["id"]
    r = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("lecture.txt", b"cell biology notes", "text/plain")},
    )
    assert r.status_code == 201
    src = r.json()
    assert src["origin"] == "upload"
    assert src["filename"] == "lecture.txt"

    sources = client.get(f"/api/spaces/{sid}/sources").json()
    assert len(sources) == 1

    detail = client.get(f"/api/spaces/{sid}").json()
    assert detail["sources"][0]["id"] == src["id"]


def test_upload_unsupported_type(client):
    sid = client.post("/api/spaces", json={"name": "Bio"}).json()["id"]
    r = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("pic.png", b"\x89PNG\x00", "image/png")},
    )
    assert r.status_code == 415


def test_delete_source(client):
    sid = client.post("/api/spaces", json={"name": "Bio"}).json()["id"]
    src = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("a.txt", b"data", "text/plain")},
    ).json()
    assert client.delete(f"/api/spaces/{sid}/sources/{src['id']}").status_code == 204
    assert client.get(f"/api/spaces/{sid}/sources").json() == []


def test_upload_sanitizes_filename(client):
    sid = client.post("/api/spaces", json={"name": "Bio"}).json()["id"]
    r = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("../../etc/passwd.txt", b"data", "text/plain")},
    )
    assert r.status_code == 201
    assert r.json()["filename"] == "passwd.txt"
    r = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": (r"C:\Users\me\notes<1>.txt", b"data", "text/plain")},
    )
    assert r.status_code == 201
    assert "/" not in r.json()["filename"] and "\\" not in r.json()["filename"]


def test_upload_too_large(client, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "max_upload_bytes", 10)
    sid = client.post("/api/spaces", json={"name": "Bio"}).json()["id"]
    r = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("big.txt", b"x" * 11, "text/plain")},
    )
    assert r.status_code == 413


def test_upload_corrupt_or_empty_file(client):
    sid = client.post("/api/spaces", json={"name": "Bio"}).json()["id"]
    corrupt = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("bad.pdf", b"not really a pdf", "application/pdf")},
    )
    assert corrupt.status_code == 422
    empty = client.post(
        f"/api/spaces/{sid}/sources",
        files={"file": ("empty.txt", b"   \n", "text/plain")},
    )
    assert empty.status_code == 422
