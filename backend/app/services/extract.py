import io

TEXT_EXTS = {
    "txt", "md", "csv", "json", "py", "js", "ts", "html", "xml", "yml", "yaml",
}


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a file's bytes. Raises ValueError if unsupported."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(
            (page.extract_text() or "") for page in reader.pages
        ).strip()

    if ext == "docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip()

    if ext == "pptx":
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        lines = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.append(shape.text_frame.text)
        return "\n".join(lines).strip()

    if ext in TEXT_EXTS:
        return data.decode("utf-8", errors="replace").strip()

    raise ValueError(f"Unsupported file type: .{ext or '(none)'}")
